"""encar_eval 오프라인 평가 결과(eval_market_price_encar.py 산출물)를 검증용 엑셀 한 장으로
묶는다 — 사람이 지표·케이스·사람판정 결과를 한 파일에서 훑을 수 있게(report.md §1 배경).

시트 7개: 요약(모드×방법 지표) · 케이스_loo · 케이스_holdout · 세대별(target_slug별 지표) ·
보정(TabPFN 분위수 커버리지·verdict 분포 대 기대치) · 사람판정(무작위 표본, 사람이 채울 빈 칸) ·
데이터설명(report.md §2 원문 그대로).
"""

import argparse
import csv
import json
import pathlib
import random

from openpyxl import Workbook
from openpyxl.styles import Font

DEFAULT_OUT_DIR = pathlib.Path("/home/whlee/workspace/bmad-encar-demo/.logs/encar_eval/20260905")
DEFAULT_REPORT_MD = DEFAULT_OUT_DIR / "report.md"
MODES = ["loo", "holdout"]
METHODS = ["tabpfn_price", "comps_median", "regression_pred"]
VERDICT_ORDER = ["저렴", "다소 저렴", "적정", "다소 높음", "높음", "(없음)"]


def _maybe_num(s):
    if s is None or s == "":
        return None
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return s


def _bold_header(ws, headers: list) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)


def load_metrics(in_dir: pathlib.Path) -> dict:
    out = {}
    for mode in MODES:
        p = in_dir / f"metrics_{mode}.json"
        if p.exists():
            out[mode] = json.loads(p.read_text(encoding="utf-8"))
    return out


def load_cases_csv(in_dir: pathlib.Path, mode: str):
    p = in_dir / f"cases_{mode}.csv"
    if not p.exists():
        return [], []
    with p.open(newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    if not rows:
        return [], []
    return rows[0], rows[1:]


def sheet_summary(wb: Workbook, metrics: dict) -> None:
    ws = wb.create_sheet("요약")
    _bold_header(ws, ["mode", "method", "n", "MdAPE", "MAPE", "hit@5%", "hit@10%", "hit@20%"])
    for mode, m in metrics.items():
        for method in METHODS:
            s = m["overall"][method]
            ws.append([mode, method, s["n"], s["mdape"], s["mape"], s["hit5"], s["hit10"], s["hit20"]])
    if not metrics:
        ws.append(["(metrics_loo.json / metrics_holdout.json 없음)"])


def sheet_cases(wb: Workbook, name: str, header: list, rows: list) -> None:
    ws = wb.create_sheet(name)
    if not header:
        ws.append([f"(cases_{name.split('_')[-1]}.csv 없음)"])
        return
    _bold_header(ws, header)
    for row in rows:
        ws.append([_maybe_num(v) for v in row])


def sheet_per_target(wb: Workbook, metrics: dict) -> None:
    ws = wb.create_sheet("세대별")
    _bold_header(ws, ["mode", "target_slug", "method", "n", "MdAPE", "MAPE", "hit@5%", "hit@10%", "hit@20%"])
    for mode, m in metrics.items():
        for slug, per_method in sorted(m["per_target"].items()):
            for method in METHODS:
                s = per_method[method]
                ws.append([mode, slug, method, s["n"], s["mdape"], s["mape"], s["hit5"], s["hit10"], s["hit20"]])


def sheet_calibration(wb: Workbook, metrics: dict) -> None:
    ws = wb.create_sheet("보정")
    _bold_header(ws, ["mode", "지표", "관측", "기대(%)"])
    for mode, m in metrics.items():
        extra = m["tabpfn_extra"]
        for key in ("coverage_q10_q90", "coverage_q25_q75", "share_basis_not_quantile",
                    "train_n_mean", "train_n_median", "elapsed_mean_s"):
            ws.append([mode, key, extra[key], None])
        n = m["n_cases"] or 1
        for verdict in VERDICT_ORDER:
            observed = extra["verdict_counts"].get(verdict, 0)
            expected_pct = extra["verdict_expected_pct"].get(verdict)
            ws.append([mode, f"verdict={verdict}", f"{observed} ({observed / n:.0%})", expected_pct])


def sheet_human_judgement(wb: Workbook, in_dir: pathlib.Path, seed: int = 20260905, k: int = 50) -> None:
    ws = wb.create_sheet("사람판정")
    header = ["encar_id", "model", "year", "mileage", "actual_price", "tabpfn_price", "verdict",
               "사람 판정(저렴·적정·높음)", "사람 근거", "일치(1/0)"]
    _bold_header(ws, header)
    wanted = ["encar_id", "model", "year", "mileage", "actual_price", "tabpfn_price", "verdict"]
    pool = []
    for mode in MODES:
        cols, rows = load_cases_csv(in_dir, mode)
        if not cols:
            continue
        idx = {c: i for i, c in enumerate(cols)}
        if not all(w in idx for w in wanted):
            continue
        for row in rows:
            pool.append([_maybe_num(row[idx[w]]) for w in wanted])
    sample = random.Random(seed).sample(pool, min(k, len(pool))) if pool else []
    for row in sample:
        ws.append(row + ["", "", ""])


def sheet_data_description(wb: Workbook, report_md_path: pathlib.Path) -> None:
    ws = wb.create_sheet("데이터설명")
    _bold_header(ws, ["report.md §2 원문(수집 조건 — 사실 기록)"])
    if not report_md_path.exists():
        ws.append([f"(파일 없음: {report_md_path})"])
        return
    lines = report_md_path.read_text(encoding="utf-8").splitlines()
    start = end = None
    for i, line in enumerate(lines):
        if line.startswith("## 2."):
            start = i
        elif start is not None and i > start and line.startswith("## "):
            end = i
            break
    if start is None:
        ws.append(["(report.md에서 '## 2.' 헤더를 찾지 못함)"])
        return
    for line in lines[start : end if end is not None else len(lines)]:
        ws.append([line])


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--in", dest="in_dir", default=str(DEFAULT_OUT_DIR / "eval"))
    ap.add_argument("--report-md", default=str(DEFAULT_REPORT_MD))
    ap.add_argument("--out", default=None, help="기본값: <in>/verification.xlsx")
    args = ap.parse_args()

    in_dir = pathlib.Path(args.in_dir)
    out_path = pathlib.Path(args.out) if args.out else in_dir / "verification.xlsx"
    metrics = load_metrics(in_dir)

    wb = Workbook()
    wb.remove(wb.active)

    sheet_summary(wb, metrics)
    for mode in MODES:
        cols, rows = load_cases_csv(in_dir, mode)
        sheet_cases(wb, f"케이스_{mode}", cols, rows)
    sheet_per_target(wb, metrics)
    sheet_calibration(wb, metrics)
    sheet_human_judgement(wb, in_dir)
    sheet_data_description(wb, pathlib.Path(args.report_md))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"작성 완료: {out_path}")
    print("시트:", wb.sheetnames)


if __name__ == "__main__":
    main()
